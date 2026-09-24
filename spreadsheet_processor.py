"""Spreadsheet input/output and one controlled recovery pass."""
import csv
import io
import logging
from collections import Counter
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, PatternFill, Side

from hollis_api import SearchError, SearchResult
from isbn_utils import normalize_isbn

LOG = logging.getLogger(__name__)
FILLS = {
    "green": PatternFill("solid", fgColor="DDF1D8"),
    "red": PatternFill("solid", fgColor="F7D9D9"),
    "yellow": PatternFill("solid", fgColor="FFF0BD"),
}
THIN = Side(style="thin", color="B7B7B7")
ALL_BORDERS = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
OUTPUT_HEADER = "Title of Book Found"


def read_input(filename, contents):
    extension = Path(filename).suffix.lower()
    if extension == ".xlsx":
        return load_workbook(io.BytesIO(contents))
    if extension == ".csv":
        try:
            rows = list(csv.reader(io.StringIO(contents.decode("utf-8-sig"))))
        except (UnicodeError, csv.Error) as exc:
            raise ValueError("CSV must be valid UTF-8.") from exc
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Books"
        for row in rows:
            sheet.append(row)
            # Treat CSV data as literal text, including values starting '='.
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        return workbook
    raise ValueError("Upload an .xlsx or UTF-8 .csv file.")


def column_options(sheet):
    from openpyxl.utils import get_column_letter
    options = {}
    for col in range(1, sheet.max_column + 1):
        header = sheet.cell(1, col).value
        if header is None and not any(sheet.cell(row, col).value is not None
                                      for row in range(2, sheet.max_row + 1)):
            continue
        label = str(header).strip() if header is not None else "(blank header)"
        options[f"{get_column_letter(col)} — {label}"] = col
    return options


def data_rows(sheet, original_max_col):
    return [row for row in range(2, sheet.max_row + 1)
            if any(sheet.cell(row, col).value not in (None, "")
                   for col in range(1, original_max_col + 1))]


def classify_rows(sheet, rows, isbn_col, client, progress=None, status=None):
    """Search each unique ISBN, then retry transient failures once.

    Persistent 429s pause the FIRST pass and queue untouched ISBNs for
    recovery. A persistent 429 in recovery ends the run. Long Retry-After
    instructions or authentication failures end it without recovery.
    """
    groups, results, outcomes = {}, {}, {}
    client.run_notice = ""
    for row in rows:
        raw = sheet.cell(row, isbn_col).value
        isbn = normalize_isbn(raw)
        if isbn is None:
            results[row] = SearchResult("white", reason="Missing or invalid ISBN.")
        else:
            groups.setdefault(isbn, []).append(row)

    def update(isbn, result):
        outcomes[isbn] = result
        for row in groups[isbn]:
            results[row] = result
        if progress:
            progress(len(results), len(rows))

    def run_pass(keys, label):
        for position, isbn in enumerate(keys, 1):
            if status:
                status(f"{label}: ISBN {position} of {len(keys)}.")
            error = None
            try:
                result = client.search_isbn(isbn)
            except SearchError as exc:
                error = exc
                result = SearchResult("white", reason=str(exc),
                                      retryable=exc.retryable, status_code=exc.status_code)
            except Exception as exc:
                # Log type only: exception text may contain request details.
                LOG.error("Unexpected search error: %s", type(exc).__name__)
                result = SearchResult("white", reason="Unexpected search error.")
            update(isbn, result)
            if error is not None and error.stop_batch:
                return error
        return None

    if progress:
        progress(len(results), len(rows))
    stopped = run_pass(list(groups), "First pass")
    can_recover = stopped is None or (stopped.retryable and not stopped.defer_run)
    pending = [isbn for isbn in groups
               if isbn not in outcomes or outcomes[isbn].retryable]
    if pending and can_recover:
        try:
            client.wait_for_second_pass()
        except SearchError as exc:
            stopped = exc
        else:
            # Deliberately call the API again. Do not reuse cached failures.
            stopped = run_pass(pending, "Recovery pass")

    if stopped is not None:
        client.run_notice = f"Run stopped: {stopped} Completed results are available below."
    for isbn in groups:
        if isbn not in outcomes:
            reason = "Not searched because the batch was stopped."
            if stopped is not None:
                reason += " " + str(stopped)
            update(isbn, SearchResult("white", reason=reason))
        elif outcomes[isbn].retryable:
            # Keep the actual failure visible; no further automatic pass.
            if stopped is not None:
                outcomes[isbn].reason += " Further automatic retries were stopped."
    if progress:
        progress(len(rows), len(rows))
    return results


def render_workbook(workbook, sheet_name, rows, results):
    sheet = workbook[sheet_name]
    original_max_col = sheet.max_column
    sheet.insert_cols(1)  # Classification already finished using original positions.
    last_col = original_max_col + 1
    header = sheet.cell(1, 1)
    header.value = OUTPUT_HEADER
    header._style = copy(sheet.cell(1, 2)._style)
    sheet.column_dimensions["A"].width = 48
    counts, preview = Counter(), []
    for row in rows:
        result = results[row]
        counts[result.status] += 1
        output = sheet.cell(row, 1)
        if result.matches and result.status in ("red", "yellow"):
            if result.status == "red":
                match = result.matches[0]
                output.value = match.title
                if match.url:
                    output.hyperlink = match.url
                    font = copy(output.font)
                    font.color, font.underline = "0563C1", "single"
                    output.font = font
            else:
                output.value = "\n".join(
                    f"{m.title} — {m.url}" if m.url else m.title for m in result.matches)
                alignment = copy(output.alignment)
                alignment.wrap_text = True
                output.alignment = alignment
            output.data_type = "s"  # API titles are text, never Excel formulas.
        for col in range(1, last_col + 1):
            cell = sheet.cell(row, col)
            if result.status in FILLS:
                cell.fill = copy(FILLS[result.status])
            cell.border = copy(ALL_BORDERS)
        preview.append({"Excel row": row, "Status": result.status,
                        "Title of Book Found": output.value or "",
                        "Review note": result.reason})
    for col in range(1, last_col + 1):
        sheet.cell(1, col).border = copy(ALL_BORDERS)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue(), counts, preview
