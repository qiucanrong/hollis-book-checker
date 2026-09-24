"""Streamlit interface for HOLLIS ISBN purchase screening."""
import hashlib
import logging
import os
from pathlib import Path

import streamlit as st

from hollis_api import HollisClient
from spreadsheet_processor import (
    classify_rows, column_options, data_rows, read_input, render_workbook,
    ImportErrorMessage, result_notes,
)

logging.basicConfig(
    level=getattr(logging, os.getenv("HOLLIS_LOG_LEVEL", "WARNING").upper(), logging.WARNING),
    format="%(asctime)s %(levelname)s %(name)s: %(message)s",
)
st.set_page_config(page_title="HOLLIS ISBN Checker", layout="wide")
st.title("HOLLIS ISBN Purchase Screening")
st.write("Upload a book list, select its ISBN and title columns, and download a checked Excel copy.")
st.caption("Green: no verified match · Red: one match · Yellow: multiple matches · Uncolored: manual review")
uploaded = st.file_uploader("Book spreadsheet", type=["xlsx", "csv", "tsv", "txt"])
if uploaded is None:
    st.stop()
contents = uploaded.getvalue()
file_hash = hashlib.sha256(contents).hexdigest()
encoding, delimiter = "Auto", "Auto"
if Path(uploaded.name).suffix.lower() != ".xlsx":
    with st.expander("Text file options (change these if the preview looks wrong)"):
        encodings = {
            "Auto": "Auto", "UTF-8": "utf-8-sig", "UTF-16 (with BOM)": "utf-16",
            "UTF-16 LE": "utf-16-le", "UTF-16 BE": "utf-16-be",
            "UTF-32 (with BOM)": "utf-32", "Windows-1252 (Western European)": "cp1252",
            "Latin-1": "latin-1", "Windows-1255 (Hebrew)": "cp1255",
            "Windows-1251 (Cyrillic)": "cp1251", "Shift JIS (Japanese)": "shift_jis",
            "GB18030 (Chinese)": "gb18030", "Mac Roman": "mac_roman",
        }
        separators = {"Auto": "Auto", "Tab": "\t", "Comma": ",", "Semicolon": ";", "Pipe": "|"}
        encoding = encodings[st.selectbox("Text encoding", list(encodings), key=f"encoding:{file_hash}")]
        delimiter = separators[st.selectbox("Column separator", list(separators), key=f"delimiter:{file_hash}")]
import_info = {}
try:
    workbook = read_input(uploaded.name, contents, encoding, delimiter, import_info)
except ImportErrorMessage as exc:
    st.error(str(exc))
    st.stop()
except Exception:
    st.error("Could not read this file. Use .xlsx or a delimited text export; check Text file options for CSV/TSV/TXT files.")
    st.stop()
for warning in import_info.get("warnings", []):
    st.warning(warning)
if import_info:
    separator_name = {"\t": "tabs", ",": "commas", ";": "semicolons", "|": "pipes"}[import_info["delimiter"]]
    st.caption(f"Imported text using {import_info['encoding']} and {separator_name}.")
sheet_name = st.selectbox("Worksheet", workbook.sheetnames, key=f"sheet:{file_hash}")
sheet = workbook[sheet_name]
header_context = (file_hash, uploaded.name, encoding, delimiter, sheet_name)
if st.session_state.get("header_context") != header_context:
    st.session_state["header_row"] = 1
    st.session_state["header_context"] = header_context
header_row = st.selectbox("Header row (column names)", list(range(1, min(5, sheet.max_row) + 1)),
                          key="header_row", help="Only rows below this row will be searched.")
with st.expander("Preview the first five rows"):
    from openpyxl.utils import get_column_letter
    st.dataframe([
        {"Row": r, **{get_column_letter(c): str(sheet.cell(r, c).value or "")
                      for c in range(1, min(sheet.max_column, 12) + 1)}}
        for r in range(1, min(sheet.max_row, 5) + 1)
    ], hide_index=True)
    if sheet.max_column > 12:
        st.caption("This preview shows columns A–L. All columns are available in the selectors below.")
options = column_options(sheet, header_row)
if len(options) < 2:
    st.error("Choose a header row with separate ISBN and title columns.")
    st.stop()
labels = list(options)
isbn_default = next((i for i, v in enumerate(labels) if "isbn" in v.lower()), 0)
title_default = next((i for i, v in enumerate(labels) if "title" in v.lower()), 0)
selection_context = f"{header_context}:{header_row}"
isbn_label = st.selectbox("ISBN column", labels, index=isbn_default, key=f"isbn:{selection_context}")
title_label = st.selectbox("Original title column", labels, index=title_default, key=f"title:{selection_context}")
environment = st.selectbox("HOLLIS environment", ["Sandbox", "Production"])
st.caption("Use the API key issued for the selected environment.")
isbn_col, title_col = options[isbn_label], options[title_label]
if isbn_col == title_col:
    st.warning("Select different ISBN and title columns.")
    st.stop()
rows = data_rows(sheet, sheet.max_column, header_row)
st.write(f"{len(rows)} nonempty book rows found.")
signature = (file_hash, uploaded.name, encoding, delimiter,
             sheet_name, header_row, isbn_col, title_col, environment)
if st.session_state.get("result_signature") != signature:
    st.session_state.pop("completed_result", None)

if st.button("Search HOLLIS", disabled=not rows):
    try:
        api_key = st.secrets["HOLLIS_API_KEY"]
        if not isinstance(api_key, str) or not api_key.strip():
            raise KeyError("HOLLIS_API_KEY")
    except Exception:
        st.error("Check HOLLIS_API_KEY in your Streamlit secrets configuration.")
        st.stop()
    st.session_state.pop("completed_result", None)
    bar, row_message = st.progress(0), st.empty()
    phase_message, request_message = st.empty(), st.empty()

    def update_progress(done, total):
        bar.progress(done / total)
        row_message.write(f"{done} of {total} rows have a result; temporary failures may be retried.")

    client = HollisClient(api_key.strip(), environment, status_callback=request_message.info)
    try:
        results = classify_rows(sheet, rows, isbn_col, client,
                                update_progress, status=phase_message.write)
    finally:
        client.close()
    notice = client.run_notice
    phase_message.empty()
    request_message.empty()
    row_message.write("Processing finished.")
    try:
        output, counts, preview = render_workbook(workbook, sheet_name, rows, results, header_row)
    except Exception:
        st.error("Searches finished, but the workbook could not be exported. The results are shown below.")
        st.dataframe([{"Excel row": r, "Status": v.status, "Notes": result_notes(v)}
                      for r, v in results.items()], hide_index=True)
        st.stop()
    filename = f"Searched_{Path(uploaded.name).stem}.xlsx"
    st.session_state["result_signature"] = signature
    st.session_state["completed_result"] = (output, filename, dict(counts), preview, notice)

completed = st.session_state.get("completed_result")
if completed and st.session_state.get("result_signature") == signature:
    output, filename, counts, preview, notice = completed
    if notice:
        st.warning(notice)
    st.subheader("Results")
    values = [("Rows", len(rows)), ("Purchase candidates", counts.get("green", 0)),
              ("Already in HOLLIS", counts.get("red", 0)),
              ("Multiple records", counts.get("yellow", 0)),
              ("Manual review", counts.get("white", 0))]
    for column, (label, value) in zip(st.columns(5), values):
        column.metric(label, value)
    st.dataframe(preview, hide_index=True, use_container_width=True)
    st.download_button("Download checked workbook", data=output, file_name=filename,
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
